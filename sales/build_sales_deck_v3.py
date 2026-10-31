"""U-ライティング 1on1営業資料 v3（台本ベースで全面刷新）

usage: python3 build_v3.py <assets_dir> <out.pptx>
"""
import sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

A, OUT = sys.argv[1], sys.argv[2]
FONT = 'Noto Sans JP'
INK = RGBColor(0x2B, 0x2B, 0x2B)
SUB = RGBColor(0x5E, 0x5E, 0x5E)
GOLD = RGBColor(0xC6, 0xA1, 0x5B)
CREAM = RGBColor(0xF7, 0xF3, 0xEA)
LINE = RGBColor(0xE5, 0xDC, 0xC8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
W, H = 12192000, 6858000
MX = 762000          # 左右マージン
CW = W - 2 * MX      # コンテンツ幅

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
page = [0]


# ------------------------------------------------------------ helpers
def font(run, size, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set('typeface', FONT)


def text(s, x, y, w, h, lines, size=14, bold=False, color=INK,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.25):
    """lines: str | list of str | list of (str, size, bold, color)."""
    tb = s.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        if isinstance(ln, str):
            ln = (ln, size, bold, color)
        t, sz, b, c = ln
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = t
        font(r, sz, b, c)
    return tb


def box(s, x, y, w, h, fill=WHITE, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1):
    b = s.shapes.add_shape(shape, Emu(x), Emu(y), Emu(w), Emu(h))
    b.shadow.inherit = False
    if fill is None:
        b.fill.background()
    else:
        b.fill.solid()
        b.fill.fore_color.rgb = fill
    if line is None:
        b.line.fill.background()
    else:
        b.line.color.rgb = line
        b.line.width = Pt(lw)
    return b


def pic(s, path, x, y, w=None, h=None, border=False):
    iw, ih = Image.open(path).size
    if w and not h:
        h = int(w * ih / iw)
    if h and not w:
        w = int(h * iw / ih)
    p = s.shapes.add_picture(path, Emu(x), Emu(y), Emu(w), Emu(h))
    if border:
        p.line.color.rgb = LINE
        p.line.width = Pt(1)
    return p


def notes(s, body):
    s.notes_slide.notes_text_frame.text = body.strip()


def new(dark=False):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = INK if dark else WHITE
    page[0] += 1
    return s


def base(label, title, speech):
    s = new()
    s.shapes.add_picture(f'{A}/topbar.png', 0, 0, Emu(W), Emu(142875))
    text(s, MX, 470000, CW, 260000, label, 11, True, GOLD)
    text(s, MX, 740000, CW, 600000, title, 26, True)
    box(s, MX, 6420000, CW, 9525, fill=LINE)
    text(s, MX, 6500000, 4000000, 200000, '© U-Writing All Rights Reserved.', 8, False, SUB)
    text(s, W - MX - 600000, 6470000, 600000, 260000, str(page[0]), 12, True, INK, PP_ALIGN.RIGHT)
    notes(s, speech)
    return s


def section(no, title, sub, speech):
    s = new(dark=True)
    text(s, MX, 2300000, 3000000, 900000, no, 54, True, GOLD)
    box(s, MX, 3250000, 900000, 38100, fill=GOLD)
    text(s, MX, 3450000, CW, 700000, title, 36, True, WHITE)
    text(s, MX, 4250000, CW, 500000, sub, 15, False, RGBColor(0xCF, 0xCF, 0xCF))
    notes(s, speech)
    return s


def tag(s, x, y, label, fill=INK, color=WHITE, size=10, w=None):
    w = w or (len(label) * 150000 + 300000)
    b = box(s, x, y, w, 300000, fill=fill)
    tf = b.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    font(r, size, True, color)
    return b


def card(s, x, y, w, h, head, body, num=None, fill=WHITE, head_size=16, body_size=12, accent=True):
    box(s, x, y, w, h, fill=fill, line=LINE)
    if accent:
        box(s, x, y, w, 50800, fill=GOLD)
    ty = y + 280000
    if num:
        text(s, x + 280000, ty, w - 560000, 300000, num, 11, True, GOLD)
        ty += 330000
    text(s, x + 280000, ty, w - 560000, 800000, head, head_size, True)
    text(s, x + 280000, ty + (head_size * 12700 * 2 + 250000 if '\n' in head else head_size * 12700 + 280000),
         w - 560000, h - (ty - y) - 400000, body, body_size, False, SUB, spacing=1.45)


def quote(s, x, y, w, h, q, who):
    box(s, x, y, w, h, fill=CREAM)
    text(s, x + 300000, y + 150000, 600000, 600000, '“', 54, True, GOLD)
    text(s, x + 300000, y + 750000, w - 600000, h - 1300000, q, 14, False, INK, spacing=1.5)
    text(s, x + 300000, y + h - 500000, w - 600000, 300000, who, 11, True, SUB, PP_ALIGN.RIGHT)


# ================================================================ slides
# 1 表紙 ---------------------------------------------------------------
s = new()
cov = Image.open(f'{A}/cover.png')
cov.crop((430, 0, cov.size[0], cov.size[1])).save(f'{A}/cover_crop.png')
pic(s, f'{A}/cover_crop.png', W - 6300000, 0, h=H)
text(s, MX, 1500000, 6000000, 300000, 'U-ライティング サービスのご案内', 13, True, SUB)
box(s, MX, 1900000, 700000, 38100, fill=GOLD)
text(s, MX, 2150000, 6200000, 1800000, ['マーケティング迷子を、', '減らす。'], 44, True, spacing=1.15)
text(s, MX, 3950000, 6000000, 400000, '薬機法を守りながら“売れる”広告づくりと集客支援', 15, False, SUB)
text(s, MX, 5300000, 6000000, 600000, [('集客プランナー / ライター', 11, True, GOLD), ('上原 龍之介', 18, True, INK)])
notes(s, """
【画面共有はまだしない】①アイスブレイクは顔を見て話す。

▼アイスブレイク（1つ選ぶ）
・紹介で会う場合：「○○さんとはどういうご縁なんですか？」（共通の知人の話）
・マッチングサービス経由：「プロフィールで△△を拝見して、気になっていたんです」
・交流会で会った場合：「先日の□□でお話しした件の続きなんですが」

▼自分から連絡した場合の理由（必ず具体的に一言）
・人をつなぐ方へ：「○○さんの周りにいる美容・健康系の事業者さんのお役に立てると思って、お会いしたかったんです」
・見込み客へ：「○○さんの（Instagram／LP／チラシ）を拝見して、△△の伝え方がもったいないと感じたんです」
　「もったいない」はOK。「間違っている」はNG。
""")

# 2 本日の流れ -----------------------------------------------------------
s = base('AGENDA', '本日の流れ', """
今日はまず○○さんのお仕事を伺って、そのあと私の話をさせてください。
私は「マーケティング迷子を減らす」をテーマにしているので、その観点でお役に立てることがないか、一緒に探せたら嬉しいです。
""")
steps = [('01', '○○様のお仕事', 'お仕事の内容や、集客の今を\n伺います'),
         ('02', '私の仕事と想い', 'なぜこの仕事をしているかを\nお話しします'),
         ('03', 'お役に立てること', 'サービスと実績を\nご紹介します'),
         ('04', '次の一歩', '今日のお話をもとに\n進め方を決めます')]
cw, gap = 2520000, 196000
for i, (n, h_, b_) in enumerate(steps):
    x = MX + i * (cw + gap)
    box(s, x, 1900000, cw, 3600000, fill=CREAM if i == 0 else WHITE, line=LINE)
    text(s, x + 280000, 2150000, cw, 700000, n, 40, True, GOLD)
    text(s, x + 280000, 3000000, cw - 400000, 500000, h_, 18, True)
    text(s, x + 280000, 3650000, cw - 400000, 1200000, b_.split('\n'), 12, False, SUB, spacing=1.5)
text(s, MX, 5700000, CW, 400000, '※ 最初は○○様のお話を中心に伺います', 12, False, SUB)

# 3 伺いたいこと（ヒアリング補助） ---------------------------------------
s = base('HEARING', 'まずは○○様のことを教えてください', """
【任意。共有しなくてもよい】ヒアリングの手元メモとして使う。
提案したくなっても我慢。相手の課題は一言でメモしておく。

▼見込み客の場合に追加（課題の深さ・お金・時期）
・その課題で、どれくらい損をしていそうですか？
・今、集客にどれくらいかけていますか？
・いつまでに解決したいですか？

▼③に移るタイミング（どちらか）
・相手の課題を要約して「そういうことですね」と同意が取れたとき
・開始から20〜25分が経ったとき
「お話を伺って、○○さんは△△のところで少し迷子になっているように感じました。私がなぜこの仕事をしているかも含めて、少しお話しさせてください。」
""")
qs = [('お仕事のこと', ['どんなお客様に、何を提供されていますか？', '今のお仕事を始めたきっかけは？']),
      ('集客のこと', ['「これで合ってるのかな」と感じていることは？', 'やってみたけど、効果が分からなかったものは？', '逆に、うまくいったと感じたのは？']),
      ('お知り合いのこと', ['普段どんな業種の方とつながっていますか？', '集客に困っていそうな方はいますか？'])]
cw, gap = 3400000, 234000
for i, (h_, items) in enumerate(qs):
    x = MX + i * (cw + gap)
    card(s, x, 1800000, cw, 4300000, h_, ['・' + q for q in items], num=f'0{i+1}', body_size=13)

# 4 セクション：私の仕事と想い ----------------------------------------------
section('01', '私の仕事と想い', 'なぜ、この仕事をしているのか', """
（②ヒアリングから移るときの一言）
お話を伺って、○○さんは△△のところで少し迷子になっているように感じました。
私がなぜこの仕事をしているかも含めて、少しお話しさせてください。
""")

# 5 自己紹介 -------------------------------------------------------------
s = base('PROFILE', '上原 龍之介', """
上原龍之介です。美容・健康分野を中心に、薬機法を守りながら“売れる”広告を作っています。
ライター歴は6年以上で、今は公式LINEやMeta広告の運用まで支援しています。
（30秒で終える。略歴は聞かれたら答える）
""")
pic(s, f'{A}/profile.jpg', MX, 1650000, w=3900000)
x0 = MX + 3900000 + 500000
text(s, x0, 1700000, 6000000, 300000, '集客プランナー / ライター', 13, True, GOLD)
text(s, x0, 2050000, 6300000, 1100000, ['美容・健康分野で、', '薬機法を守りながら“売れる”広告を作っています。'], 20, True, spacing=1.35)
facts = [('6年以上', 'ライター歴'), ('LP・LINE・広告', '企画から運用まで'), ('薬機法・景表法', '配慮した表現設計')]
fw = 2000000
for i, (big, small) in enumerate(facts):
    x = x0 + i * (fw + 150000)
    box(s, x, 3450000, fw, 1050000, fill=CREAM)
    text(s, x, 3600000, fw, 450000, big, 16, True, INK, PP_ALIGN.CENTER)
    text(s, x, 4080000, fw, 300000, small, 11, False, SUB, PP_ALIGN.CENTER)
text(s, x0, 4800000, 6300000, 1300000, [
    ('略歴', 11, True, GOLD),
    ('ゲーム制作の専門学校を卒業 → ゲームアプリのメールサポート（2年）', 11, False, SUB),
    ('→ 建築積算（3.5年） → フリーランスライターとして独立', 11, False, SUB),
    ('→ 現在：ライター兼集客プランナー', 11, False, SUB)], spacing=1.5)

# 6 よく聞く声 -----------------------------------------------------------
s = base('WHY I DO THIS', 'いろんな方とお話しする中で、本当によく聞く話があります', """
私がこの仕事を続けている理由を、少しだけお話しさせてください。
いろんな事業者さんとお話ししてきた中で、本当によく聞く話があるんです。
「何十万円もかけてホームページを作ったけど、問い合わせが来ない」
「高い講座やコンサルを受けたけど、成果につながらなかった」。
""")
quote(s, MX, 1900000, 5180000, 3600000, ['何十万円もかけて', 'ホームページを作ったけど、', '問い合わせが来ない。'], '― Webページを作った事業者さん')
quote(s, MX + 5488000, 1900000, 5180000, 3600000, ['高い講座やコンサルを', '受けたけど、', '成果につながらなかった。'], '― 講座・コンサルを受けた事業者さん')

# 7 自力 vs 専門家 ---------------------------------------------------------
s = base('WHY I DO THIS', 'それでも、皆さん諦めない。でも――', """
それでも、皆さん諦めないんですよね。今度は自力で、SNSやチラシで集客を頑張ろうとする。
でも、自力で結果を出すには、どうしても時間がかかります。
専門家の力を借りれば早いけれど、またお金がかかる。

【注意】「支援業者は高すぎる」という言い方は絶対にしない。同席者や紹介先にコンサル・制作会社がいる。
""")
for i, (h_, b_, pro) in enumerate([('自力で頑張る', '手探りなので、結果が出るまで', '時間がかかる'),
                                     ('専門家に頼む', '早く結果は出せるけれど、また', 'お金がかかる')]):
    x = MX + i * 5488000
    box(s, x, 1900000, 5180000, 3300000, fill=WHITE, line=LINE)
    text(s, x, 2250000, 5180000, 500000, h_, 22, True, INK, PP_ALIGN.CENTER)
    box(s, x + 2290000, 2950000, 600000, 25400, fill=GOLD)
    text(s, x, 3300000, 5180000, 400000, b_, 14, False, SUB, PP_ALIGN.CENTER)
    text(s, x, 3800000, 5180000, 700000, pro, 30, True, GOLD, PP_ALIGN.CENTER)
text(s, MX, 5500000, CW, 500000, 'どちらを選んでも、壁にぶつかってしまう。', 18, True, INK, PP_ALIGN.CENTER)

# 8 マーケティング迷子 -----------------------------------------------------
s = new(dark=True)
text(s, MX, 1700000, CW, 400000, '頑張っているのに、何にお金と時間を使えばいいか分からない。', 20, False, RGBColor(0xDD, 0xDD, 0xDD), PP_ALIGN.CENTER)
text(s, MX, 2550000, CW, 500000, '私はこれを', 18, False, RGBColor(0xDD, 0xDD, 0xDD), PP_ALIGN.CENTER)
text(s, MX, 3100000, CW, 1100000, '「マーケティング迷子」', 54, True, GOLD, PP_ALIGN.CENTER)
text(s, MX, 4350000, CW, 500000, 'と呼んでいます。', 18, False, RGBColor(0xDD, 0xDD, 0xDD), PP_ALIGN.CENTER)
notes(s, """
頑張っているのに、何にお金と時間を使えばいいか分からなくなっている。
私はこれを「マーケティング迷子」と呼んでいます。
（ここで一拍置く）
""")

# 9 広告は投資 -------------------------------------------------------------
s = base('MY BELIEF', '広告は「投資」です', """
Webページやチラシなどの広告は、一種の「投資」だと思っています。
投資の大原則は、いきなり大金をつぎ込まずに、まず少額から試すこと。
だから私は、手の届きやすい価格で、必要なところから始められるようにしています。
""")
text(s, MX, 1900000, CW, 1300000, ['投資の大原則は、', 'いきなり大金をつぎ込まず、まずは少額から試すこと。'], 24, True, INK, PP_ALIGN.CENTER, spacing=1.4)
row = [('まず小さく試す', '相談1回・LP冒頭だけの改善など、\n必要なところから'), ('反応を見て直す', '数字を見て、\n何を直すべきかを判断'), ('当たりに投資する', '成果が見えたものに、\n予算を寄せていく')]
cw, gap = 3400000, 234000
for i, (h_, b_) in enumerate(row):
    x = MX + i * (cw + gap)
    box(s, x, 3700000, cw, 2200000, fill=CREAM)
    text(s, x, 3950000, cw, 400000, f'STEP {i+1}', 11, True, GOLD, PP_ALIGN.CENTER)
    text(s, x, 4300000, cw, 500000, h_, 18, True, INK, PP_ALIGN.CENTER)
    text(s, x, 4900000, cw, 800000, b_.split('\n'), 12, False, SUB, PP_ALIGN.CENTER, spacing=1.5)

# 10 私の約束 --------------------------------------------------------------
s = base('MY PROMISE', 'だから私は、こう決めています', """
だから私は、本当に必要なものだけを作ると決めています。
ご相談を受けても、作らなくていいと思ったら正直にそう言いますし、
まずは1回の相談だけ、という使い方もできるようにしています。誇大広告もしません。
マーケティング迷子を一人でも減らしたい。それが、私の仕事の理由です。

→ この直後に②のヒアリングにつなげる：
「先ほど○○さんが△△とおっしゃっていたのが、まさにこの話だと思っていて――」
""")
promises = [('本当に必要なものだけを作る', 'ご依頼の前に、それが本当に必要かを一緒に考えます。'),
            ('作らなくていいなら、正直に言う', '不要だと判断したものは、おすすめしません。'),
            ('1回の相談から使える', '契約の縛りはなく、必要なときだけご利用いただけます。'),
            ('誇大広告はしない', '薬機法・景表法を守り、消費者とお客様が笑顔でつながる広告を作ります。')]
cw, ch = 5180000, 1850000
for i, (h_, b_) in enumerate(promises):
    x = MX + (i % 2) * (cw + 308000)
    y = 1750000 + (i // 2) * (ch + 200000)
    box(s, x, y, cw, ch, fill=WHITE, line=LINE)
    box(s, x, y, 76200, ch, fill=GOLD)
    text(s, x + 350000, y + 300000, 800000, 500000, f'0{i+1}', 24, True, GOLD)
    text(s, x + 1100000, y + 330000, cw - 1350000, 500000, h_, 17, True)
    text(s, x + 1100000, y + 900000, cw - 1350000, 800000, b_, 12, False, SUB, spacing=1.5)

# 11 セクション：お役に立てること ---------------------------------------------
section('02', 'お役に立てること', '2つのサービスと、これまでの実績', """
先ほど○○さんが△△とおっしゃっていたのが、まさにこの話だと思っていて――

▼ここで分岐する
・相手自身に集客の課題がある → 契約を目指す：サービス（12〜15）→ 事例1件（16〜19） → 実績 → 料金
・課題はないが人脈がある → 紹介を目指す：お客様の声・実績（20〜21）→ ご紹介のお願い（25〜）
・両方ある → 契約が先、紹介は締めで聞く
・どちらもない → 関係づくり。こちらから紹介できる人を出す
""")

# 12 2つのサービス ----------------------------------------------------------
s = base('SERVICE', 'あなたに合わせた、2つの支援', """
私のサービスは大きく2つです。
手を動かす時間がない方には、私が代わりに作る「カワリミ」。
自分で動きたいけど方向性に迷っている方には、横で道筋を示す「ミチスジ」です。
○○さんのお話だと、△△なので、【カワリミ／ミチスジ】が合いそうだと思っています。

【サービスは1つに絞って言い切る】
・制作に手が回らない → カワリミ
・何が悪いか分からない／自分でやりたい → ミチスジ
""")
svc = [('icon_kawarimi.png', '制作代行サービス', 'カワリミ', '時間がなく、制作まで手が回らない方へ',
        ['LP・HP制作／ライティング', '公式LINE構築・運用', 'Meta広告運用', 'チラシ・営業資料など'], '単発 5,000円〜', INK),
       ('icon_michisuji.png', '集客顧問サービス', 'ミチスジ', '自分で動きたいが、方向性に迷っている方へ',
        ['月最大4時間のZoom相談', '制作物への添削・アドバイス', '集客施策の分析・仮説検証', '1回だけの相談もOK'], '月額 49,800円〜', GOLD)]
cw = 5180000
for i, (ic, kind, name, who, items, price, col) in enumerate(svc):
    x = MX + i * (cw + 308000)
    box(s, x, 1650000, cw, 4500000, fill=WHITE, line=LINE)
    box(s, x, 1650000, cw, 1000000, fill=col)
    text(s, x + 350000, 1780000, 3500000, 300000, kind, 12, True, WHITE)
    text(s, x + 350000, 2080000, 3500000, 500000, name, 26, True, WHITE)
    pic(s, f'{A}/{ic}', x + cw - 1250000, 1450000, w=1100000)
    text(s, x + 350000, 2900000, cw - 700000, 400000, who, 15, True)
    text(s, x + 350000, 3450000, cw - 700000, 1600000, ['✓ ' + t for t in items], 13, False, SUB, spacing=1.6)
    box(s, x + 350000, 5250000, cw - 700000, 9525, fill=LINE)
    text(s, x + 350000, 5400000, cw - 700000, 500000, price, 20, True, GOLD, PP_ALIGN.RIGHT)

# 13 カワリミ ---------------------------------------------------------------
s = base('KAWARIMI', '制作代行「カワリミ」― あなたの代わりに、作って運用します', """
カワリミは、集客に必要な広告づくりを私が代わりにやるサービスです。
企画・構成から制作・運用まで丸投げいただけます。お願いするのは、必要な情報と素材のご提供だけです。
（相手に関係のあるメニューだけ指差す。全部読み上げない）
""")
menu = [('Webページ制作（LP・HP）', '100,000円〜', '構成・ライティングから実装まで'),
        ('LPライティング', '80,000円〜', '構成・原稿・ワイヤー作成'),
        ('公式LINE構築・運用', '初月50,000円〜\n月額30,000円〜', '配信代行を含む'),
        ('Meta広告運用', '初期30,000円〜\n月額 広告費の20%', 'Facebook・Instagram'),
        ('チラシ制作', '片面30,000円〜\n両面50,000円〜', ''),
        ('スピーチ原稿・営業資料', '5,000円〜\n50,000円〜', '30秒スピーチ原稿は5,000円')]
cw, ch = 3400000, 1950000
for i, (h_, price, sub) in enumerate(menu):
    x = MX + (i % 3) * (cw + 234000)
    y = 1700000 + (i // 3) * (ch + 200000)
    box(s, x, y, cw, ch, fill=WHITE, line=LINE)
    box(s, x, y, cw, 50800, fill=INK)
    text(s, x + 250000, y + 250000, cw - 500000, 400000, h_, 15, True)
    text(s, x + 250000, y + 720000, cw - 500000, 700000, price.split('\n'), 15, True, GOLD, spacing=1.2)
    if sub:
        text(s, x + 250000, y + ch - 420000, cw - 500000, 300000, sub, 10, False, SUB)
text(s, MX, 6050000, CW, 300000, '※ 価格は税込・目安です。内容やボリュームに応じて、事前にお見積りします。', 10, False, SUB)

# 14 ミチスジ ---------------------------------------------------------------
s = base('MICHISUJI', '集客顧問「ミチスジ」― 迷子にならない道筋を、一緒に', """
ミチスジは、ご自身で動きたい方の横で、道筋を示し続けるサービスです。
いきなり顧問でなくても大丈夫です。まずは1回の相談から、必要なら提案書、続けるなら顧問、と段階的に使えます。
顧問中は、制作代行が20%引きになります。
""")
plans = [('相談プラン', '14,800円〜', '1回限り', ['最大2時間のZoom相談', '現状の広告物を添削', 'まず何から手をつけるか、優先順位をご提案'], 'まず1回、現状を整理したい方'),
         ('提案プラン', '49,800円〜', '単発', ['現状分析・競合リサーチ', '具体的な施策の提案書を作成', '提案プレゼン・質疑応答'], '集客の方向性を資料で欲しい方'),
         ('顧問プラン', '49,800円〜', '月額', ['月最大4時間のZoom相談', '制作物への添削・アドバイス', '施策の分析・仮説検証', 'カワリミ 20%OFF'], '継続して伴走してほしい方')]
cw = 3400000
for i, (name, price, unit, items, who) in enumerate(plans):
    x = MX + i * (cw + 234000)
    y = 2300000 - i * 250000
    hh = 3850000 + i * 250000
    box(s, x, y, cw, hh, fill=CREAM if i < 2 else WHITE, line=GOLD if i == 2 else LINE, lw=1.5 if i == 2 else 1)
    text(s, x + 250000, y + 230000, cw - 500000, 300000, f'STEP {i+1}', 11, True, GOLD)
    text(s, x + 250000, y + 520000, cw - 500000, 450000, name, 20, True)
    text(s, x + 250000, y + 1050000, cw - 500000, 500000, [(f'{unit}  ', 11, False, SUB)], 11)
    text(s, x + 900000, y + 960000, cw - 1150000, 500000, price, 20, True, GOLD)
    text(s, x + 250000, y + 1600000, cw - 500000, 1700000, ['・' + t for t in items], 12, False, INK, spacing=1.5)
    text(s, x + 250000, y + hh - 450000, cw - 500000, 300000, '→ ' + who, 11, True, SUB)
text(s, MX, 1600000, 5000000, 400000, '小さく始めて、必要なら次へ。', 14, True, INK)

# 15 選ばれる理由 ------------------------------------------------------------
s = base('OUR STRENGTHS', '選んでいただいている3つの理由', """
選んでいただく理由は3つです。
薬機法を守りながら“売れる”表現を作れること。
作って終わりではなく、「誰に・何を・どう届けるか」から一緒に見直すこと。
そして、契約の縛りがなく、必要なときだけ使えることです。
""")
reasons = [('守りと攻めの\n薬機法対応', '法律を守りながら、単なる言い換えではない「ベネフィット」「体験」を言語化し、競合と差別化します。'),
           ('売れない“本当の原因”を\n一緒に見つける', '文章を直すだけでなく、「誰に・何を・どう届けるか」から見直し、反応が出る広告を作ります。'),
           ('1ヶ月単位で\n気軽に相談できる', '契約の縛りはありません。必要なタイミング・期間だけご利用いただけます。')]
cw = 3400000
for i, (h_, b_) in enumerate(reasons):
    x = MX + i * (cw + 234000)
    card(s, x, 1750000, cw, 4300000, h_, b_, num=f'REASON 0{i+1}', head_size=18, body_size=13)

# 16-19 事例 ---------------------------------------------------------------
cases = [
    ('CASE 01', 'クラファン｜ページ企画・ライティング担当', '機能性寝具（羽毛ふとん）',
     '薬機法上、「快眠」「疲労回復」などの効果を直接うたえず、魅力の伝え方に悩んでいた。',
     '法に触れない範囲で「睡眠体験」「朝の目覚め」の変化に焦点を当てたストーリーを設計。写真撮影のディレクションも実施。',
     '279%', '目標30万円に対し、公開1ヶ月で83.9万円', 'case1.png',
     '“快眠”と言えない寝具を、“朝の目覚め”の体験として伝えて、目標の2.8倍になりました。'),
    ('CASE 02', 'クラファン｜ページ企画・ライティング担当', '健康雑貨（ツボ押しシート）',
     'ニッチな商品で、短期間で認知を広げて支援を集める必要があった。',
     'ターゲットの「痛み」「悩み」に寄り添う構成案を作成。利用シーンが浮かぶ撮影ディレクションで、共感を呼ぶページに。',
     '219万円', '公開3日で約50万円、最終的に目標比109%', 'case2.png',
     '公開3日で約50万円、最終的に219万円まで伸びました。'),
    ('CASE 03', '広告運用・LP改善・LINE配信', 'オンライン講座',
     '広告からの流入はあるが、LPでの離脱が多く成約率が伸びない。月商も横ばいが続いていた。',
     'Meta広告の運用に加え、LPの構成とデザインを改修。LINEのステップ配信も見直し、教育から成約までの導線を最適化。',
     '最大150万円', '月平均50万円の売上から改善', None,
     '広告・LP・LINEをつなぎ直して、月平均50万円の売上が、最大で150万円まで伸びました。'),
    ('CASE 04', '採用コンテンツ制作', '建設会社（防水・塗装）の採用',
     '当初の依頼はパンフレット制作。ヒアリングすると、HPの情報不足による採用のミスマッチが本当の課題だった。',
     'パンフレット制作は中止を提案。代わりに、求める人物像と代表の想いを伝える記事をHPに掲載。',
     '1週間', '記事公開から1週間以内に、新卒1名が入社決定', 'case4.png',
     'パンフレットのご依頼でしたが、本当に必要なのは別のものだと分かって、記事に切り替えました。公開1週間で内定が出ています。\n（「必要なものだけを作る」の具体例として話せる）'),
]
for no, tg, title, prob, act, big, small, img, line in cases:
    s = base(f'{no}  導入事例', title, f"""
実際の例を一つだけお見せしますね。
{line}

【事例は相手に近いものを1件だけ】
・美容・健康の物販 → CASE 01 / 02
・講座・スクール・コーチ → CASE 03
・採用に悩む会社 → CASE 04
クラファン（01・02）は「ページ企画・ライティング担当」。聞かれたら担当範囲を正直に答える。
""")
    tag(s, MX, 1450000, tg, w=len(tg) * 140000 + 300000)
    lw_ = 5300000
    text(s, MX, 1950000, lw_, 300000, '課題', 11, True, GOLD)
    text(s, MX, 2250000, lw_, 900000, prob, 12, False, INK, spacing=1.5)
    text(s, MX, 3200000, lw_, 300000, 'やったこと', 11, True, GOLD)
    text(s, MX, 3500000, lw_, 1000000, act, 12, False, INK, spacing=1.5)
    box(s, MX, 4700000, lw_, 1500000, fill=CREAM)
    box(s, MX, 4700000, 76200, 1500000, fill=GOLD)
    text(s, MX + 300000, 4850000, 1000000, 300000, '成果', 11, True, GOLD)
    text(s, MX + 300000, 5100000, lw_ - 500000, 650000, big, 36, True, INK)
    text(s, MX + 300000, 5800000, lw_ - 500000, 300000, small, 12, False, SUB)
    rx, rw = MX + lw_ + 400000, CW - lw_ - 400000
    if img:
        iw, ih = Image.open(f'{A}/{img}').size
        h_ = min(4600000, int(rw * ih / iw))
        w_ = int(h_ * iw / ih)
        pic(s, f'{A}/{img}', rx + (rw - w_) // 2, 1500000 + (4700000 - h_) // 2, w=w_, border=True)
    else:
        for k, (st, desc) in enumerate([('Meta広告運用', '見込み客を集める'), ('LP改修', '離脱を減らし、成約率を上げる'), ('LINEステップ配信', '教育から成約までの導線を整える')]):
            y = 1600000 + k * 1250000
            box(s, rx, y, rw, 1000000, fill=WHITE, line=LINE)
            text(s, rx + 250000, y + 180000, 700000, 600000, f'0{k+1}', 26, True, GOLD)
            text(s, rx + 1000000, y + 200000, rw - 1200000, 400000, st, 16, True)
            text(s, rx + 1000000, y + 580000, rw - 1200000, 300000, desc, 12, False, SUB)
        tri = box(s, rx + rw // 2 - 250000, 5330000, 500000, 250000, fill=GOLD, shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
        tri.rotation = 180
        text(s, rx, 5650000, rw, 450000, '月の売上 最大150万円', 18, True, INK, PP_ALIGN.CENTER)

# 20 お客様の声 --------------------------------------------------------------
s = base('VOICE', 'お客様の声', """
お客様からは、こんな声をいただいています。
特に多いのが「ヒアリングが丁寧」という声で、今日のように最初にしっかりお話を伺うのは、そのためなんです。
""")
voices = [('提案力の高さ', 'こちらの要望も丁寧に聞いてくださり、的確な提案をいただけました。リーズナブルな価格にもかかわらず、しっかりとした内容でとても助かりました。', '健康装具の販売者 Sさん'),
          ('ヒアリング力の高さ', 'これまでの広告課題やお客様の反応を丁寧にヒアリングしてくださり、納得感のある広告に仕上げていただけました。次回もぜひお願いしたいと思える方です。', '脱毛・ホワイトニングサロン Yさん'),
          ('薬機法も対策可', '薬機法が関わる案件でも、安心して上原さんにすべてお任せできました。丁寧な対応と、高い完成度の原稿にとても満足しています。', 'デザイン制作会社 Hさん')]
cw = 3400000
for i, (h_, q, who) in enumerate(voices):
    x = MX + i * (cw + 234000)
    box(s, x, 1700000, cw, 4400000, fill=CREAM)
    tag(s, x + 250000, 1950000, h_, fill=GOLD, w=len(h_) * 170000 + 300000)
    text(s, x + 250000, 2450000, 500000, 600000, '“', 48, True, GOLD)
    text(s, x + 250000, 3050000, cw - 500000, 2300000, q, 13, False, INK, spacing=1.6)
    text(s, x + 250000, 5550000, cw - 500000, 300000, who, 11, True, SUB, PP_ALIGN.RIGHT)

# 21 その他の実績 ------------------------------------------------------------
s = base('WORKS', 'その他の実績', """
他にも、こんな実績があります。全部はお話ししませんが、○○さんに近いところだと――
（相手の業種に近いものを1〜2件だけ指差す。読み上げない）
・美容・化粧品・サプリ → 化粧品のLINEシナリオ「257通中44件成約、成約率17.12%」
・講座・セミナー → セミナー文章改修「0人→5人」、Meta広告「1人約14,000円で集客」
・採用 → 採用記事「公開1週間で新卒1人入社」
・物販 → クラファン3件
""")
groups = [('LINE・広告', [('17.12%', '化粧品の販促LINEシナリオ', '257通中44件が成約・単価2,000円'),
                        ('約14,000円', '講座の集客にMeta広告を活用', '1人あたりの集客単価'),
                        ('約6,750円', 'セールス代行LPのライティング', '約54,000円で8件のLINE登録')]),
          ('文章・記事', [('0人→5人', 'コーチングのセミナー文章を改修', '集客数'),
                        ('新卒1人', '防水・塗装会社の採用記事', '公開1週間以内に入社決定'),
                        ('60→70サイト', '転職サービスのPR記事を企画', '記事の平均転載数')]),
          ('クラファン', [('約30万円', 'PCガジェット', '公開3日以内（ライティング担当）'),
                       ('2件', '機能性寝具・健康雑貨', 'CASE 01・02でご紹介')])]
cw = 3400000
for i, (g, items) in enumerate(groups):
    x = MX + i * (cw + 234000)
    tag(s, x, 1600000, g, fill=INK, w=cw)
    for k, (big, what, sub) in enumerate(items):
        y = 2050000 + k * 1380000
        box(s, x, y, cw, 1250000, fill=WHITE, line=LINE)
        text(s, x + 220000, y + 150000, cw - 440000, 450000, big, 20, True, GOLD)
        text(s, x + 220000, y + 600000, cw - 440000, 300000, what, 12, True)
        text(s, x + 220000, y + 900000, cw - 440000, 300000, sub, 10, False, SUB)
text(s, MX, 6150000, CW, 250000, '※ サービスサイト（https://www.u-promotion.com/）掲載の実績より', 9, False, SUB)

# 22 過去の成果物例 ----------------------------------------------------------
s = base('PORTFOLIO', '過去の成果物例', """
実際に作ったページだと、こういったものがあります。ECサイト、LP、HPですね。
他にも公開していない実績があるので、近い事例があれば個別にお見せします。
（相手がLP・HP制作に興味がある場合だけ丁寧に。そうでなければ飛ばしてよい）
""")
items = [('ECサイト制作', 'RING A BELL OKINAWA様', '構成案・ライティング'),
         ('LP制作', 'erellaオンラインセールス代行様', '構成案・ライティング'),
         ('HP制作', '株式会社t.style様', 'ライティング')]
cw = 3400000
for k, (kind, client, scope) in enumerate(items):
    x = MX + k * (cw + 234000)
    ph = box(s, x, 1700000, cw, 2600000, fill=CREAM, line=GOLD)
    ph.line.dash_style = 4
    tf = ph.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = f'【{kind}】のスクリーンショットを貼る'
    font(r, 11, False, RGBColor(0x99, 0x88, 0x66))
    tag(s, x, 4500000, kind, fill=INK)
    text(s, x, 4950000, cw, 400000, client, 14, True)
    text(s, x, 5350000, cw, 300000, f'対応範囲：{scope}', 11, False, SUB)
text(s, MX, 5900000, CW, 300000, 'その他、非公開の実績もございます。ご希望があれば個別にお見せします。', 12, False, SUB, PP_ALIGN.CENTER)

# 23 料金一覧 ---------------------------------------------------------------
s = base('PRICE', '料金一覧', """
料金はこちらです。今日のお話だと、【○○プラン】が近いと思います。
正式な金額は、伺った内容をもとにお見積りしますね。
（全プランを説明しない。勧めるプランの行だけ指差す）
""")
rows = [('制作代行「カワリミ」', '広告物の制作・運用を代行（LP・HP、LINE、Meta広告、チラシなど）', '単発 5,000円〜'),
        ('集客顧問「ミチスジ」顧問プラン', '月最大4時間のZoom相談、添削、施策の分析・仮説検証', '月額 49,800円〜'),
        ('集客顧問「ミチスジ」相談プラン', '最大2時間のZoom相談で、添削と対応方針をご提案', '1回 14,800円〜'),
        ('集客顧問「ミチスジ」提案プラン', '現状分析と、必要な施策をまとめた提案書の作成・プレゼン', '単発 49,800円〜')]
tbl = s.shapes.add_table(len(rows) + 1, 3, Emu(MX), Emu(1700000), Emu(CW), Emu(3600000)).table
for j, wd in enumerate((3300000, 5100000, 2268000)):
    tbl.columns[j].width = Emu(wd)
for i, row_ in enumerate([('プラン', '内容', '価格（税込）')] + rows):
    for j, v in enumerate(row_):
        c = tbl.cell(i, j)
        c.fill.solid()
        c.fill.fore_color.rgb = INK if i == 0 else (CREAM if i % 2 == 0 else WHITE)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.margin_left = c.margin_right = Emu(180000)
        p = c.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT if (j == 2 and i) else PP_ALIGN.LEFT
        r = p.add_run()
        r.text = v
        font(r, 12 if i == 0 else (14 if j != 1 else 12), i == 0 or j != 1,
             WHITE if i == 0 else (GOLD if j == 2 else INK))
box(s, MX, 5550000, CW, 600000, fill=CREAM)
text(s, MX, 5550000, CW, 600000, '「ミチスジ」顧問プランの契約中は、制作代行「カワリミ」の費用が 20% OFF になります。', 13, True, INK, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

# 24 次の一歩 ---------------------------------------------------------------
s = base('NEXT STEP', '次の一歩', """
【必ず具体的な行動で終える。「またよろしくお願いします」で終わらせない】

▼契約を目指す場合：その場で提案・見積もりを送る日を決める（無料）
今日伺った内容で、必要なものだけのご提案とお見積りをお送りします。○日までにお送りしてよいですか？
・まだ迷っている相手には、相談プラン（14,800円〜）を選択肢として出す。

▼紹介を目指す場合 → 次のセクション（ご紹介のお願い）へ
""")
text(s, MX, 1600000, CW, 500000, '本日のお話をもとに、必要なものだけをご提案します。', 18, True)
nexts = [('おすすめ', 'ご提案・お見積り', '無料', '今日伺った課題に対して、制作・顧問のご提案と見積りをお送りします。'),
         ('まだ迷っている方', '相談プラン', '14,800円〜', '最大2時間のZoom相談で、現状と優先順位を一緒に整理します。'),
         ('お知り合いに', 'ご紹介', '', '集客や広告でお困りの方がいれば、ぜひおつなぎください。')]
cw = 3400000
for i, (lab, h_, price, b_) in enumerate(nexts):
    x = MX + i * (cw + 234000)
    box(s, x, 2350000, cw, 3600000, fill=CREAM if i == 0 else WHITE, line=GOLD if i == 0 else LINE, lw=2 if i == 0 else 1)
    tag(s, x + 250000, 2600000, lab, fill=GOLD if i == 0 else INK, w=len(lab) * 170000 + 300000)
    text(s, x + 250000, 3150000, cw - 500000, 500000, f'0{i+1}  {h_}', 20, True)
    if price:
        text(s, x + 250000, 3750000, cw - 500000, 500000, price, 22, True, GOLD)
    text(s, x + 250000, 4450000, cw - 500000, 1300000, b_, 13, False, SUB, spacing=1.5)

# 25 セクション：ご紹介のお願い ------------------------------------------------
section('03', 'ご紹介のお願い', 'マーケティング迷子を、一緒に減らしてください', """
（紹介を目指す場合、または契約の話のあとに）
最後に一つ、お願いがあります。
""")

# 26 こんな声を聞いたら -----------------------------------------------------
s = base('REFERRAL KEYWORDS', 'こんな声を聞いたら、思い出してください', """
もし周りで、「チラシやLPを作ったけど反応がない」「この表現、薬機法的に大丈夫？」「集客まで手が回らない」
と言っている方がいたら、私を思い出してください。合言葉は「マーケティング迷子」です。
""")
kws = [('「チラシやLPを作ったけど、反応がない」', '作った後に成果が出ていない方。原因の診断から改善までお手伝いします。'),
       ('「この表現、薬機法的に大丈夫？」', '美容・健康分野で表現に悩む方。守りながら“売れる”表現を作ります。'),
       ('「集客まで手が回らない」', '本業が忙しい方、何から手をつけるか迷っている方。代行と顧問で支えます。')]
for i, (q, b_) in enumerate(kws):
    y = 1650000 + i * 1300000
    box(s, MX, y, CW, 1150000, fill=CREAM)
    box(s, MX, y, 76200, 1150000, fill=GOLD)
    text(s, MX + 350000, y + 200000, 5700000, 800000, q, 18, True, INK, anchor=MSO_ANCHOR.MIDDLE)
    text(s, MX + 6200000, y + 200000, CW - 6500000, 800000, b_, 12, False, SUB, anchor=MSO_ANCHOR.MIDDLE, spacing=1.5)
text(s, MX, 5650000, CW, 500000, [('合言葉は「マーケティング迷子」', 20, True, GOLD)], align=PP_ALIGN.CENTER)

# 27 ご紹介いただきたい方 ----------------------------------------------------
s = base('WHO TO INTRODUCE', 'ご紹介いただきたい方', """
特に、美容サロン、整体院、医療クリニック、健康食品を扱っている事業者さんとは相性がいいです。
協業できるパートナーさんも探しています。
""")
cols = [('お客様として', '美容・健康分野で集客に悩む事業者', ['美容サロン・脱毛サロン', '整体院・治療院', '医療クリニック', '健康食品・化粧品・健康雑貨の販売者', '講座・スクール・コーチ']),
        ('協業パートナーとして', 'ライティング・薬機法対応を任せたい方', ['Web制作会社・デザイン会社', '経営・財務コンサルタント', '医療・美容機器ディーラー', 'レンタルスペース・店舗物件の不動産仲介', '経営者コミュニティ・マッチングサービス運営'])]
cw = 5180000
for i, (h_, sub, items) in enumerate(cols):
    x = MX + i * (cw + 308000)
    box(s, x, 1650000, cw, 4450000, fill=WHITE, line=LINE)
    box(s, x, 1650000, cw, 50800, fill=GOLD if i == 0 else INK)
    text(s, x + 350000, 1950000, cw - 700000, 450000, h_, 20, True)
    text(s, x + 350000, 2450000, cw - 700000, 300000, sub, 12, False, SUB)
    text(s, x + 350000, 3000000, cw - 700000, 2900000, ['● ' + t for t in items], 14, False, INK, spacing=1.7)

# 28 紹介パートナー制度 ------------------------------------------------------
s = base('REFERRAL PROGRAM', '紹介パートナー制度', """
ご紹介がご契約につながった場合は、紹介料をお支払いしています。
顧問プランなら契約が続く限り毎月20%、制作代行なら10万円以上のご契約で10%です。
商談やご提案はすべて私が行うので、つないでいただくだけで大丈夫です。
""")
fees = [('集客顧問「ミチスジ」顧問プラン', '20%', '契約金額の20%を毎月', '契約が続く限り、着金確認後にお支払い'),
        ('制作代行「カワリミ」', '10%', '契約金額の10%', '契約金額10万円以上の場合。着金確認後にお支払い')]
for i, (h_, pct, a_, b_) in enumerate(fees):
    y = 1700000 + i * 1650000
    box(s, MX, y, 5300000, 1450000, fill=CREAM)
    text(s, MX + 300000, y + 220000, 4800000, 350000, h_, 13, True)
    text(s, MX + 300000, y + 600000, 1700000, 750000, pct, 40, True, GOLD)
    text(s, MX + 2000000, y + 680000, 3200000, 300000, a_, 13, True)
    text(s, MX + 2000000, y + 1000000, 3200000, 300000, b_, 10, False, SUB)
flow = [('ご紹介・おつなぎ', 'お困りの方をご紹介ください'), ('商談・ご提案', 'ヒアリングと提案はすべて私が行います'),
        ('ご契約・業務開始', 'お客様と直接契約します'), ('紹介料のお支払い', '入金確認後にお支払いします')]
fx = MX + 5700000
for k, (h_, b_) in enumerate(flow):
    y = 1700000 + k * 850000
    box(s, fx, y, CW - 5700000, 720000, fill=WHITE, line=LINE)
    o = box(s, fx + 200000, y + 160000, 400000, 400000, fill=INK, shape=MSO_SHAPE.OVAL)
    text(s, fx + 200000, y + 160000, 400000, 400000, str(k + 1), 12, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(s, fx + 800000, y + 120000, 3800000, 300000, h_, 13, True)
    text(s, fx + 800000, y + 400000, 3800000, 300000, b_, 10, False, SUB)
text(s, MX, 5250000, CW, 700000, ['メリット：自社サービスの付加価値として提案できる／リスクなく新たな収益源になる／専門的な薬機法対応を任せられる'], 11, False, SUB)

# 29 クロージング ------------------------------------------------------------
s = new(dark=True)
text(s, MX, 1500000, CW, 500000, '本日はありがとうございました', 30, True, WHITE, PP_ALIGN.CENTER)
text(s, MX, 2400000, CW, 500000, '最後に教えてください', 14, False, RGBColor(0xCF, 0xCF, 0xCF), PP_ALIGN.CENTER)
text(s, MX, 2850000, CW, 600000, '私から○○様にご紹介できるのは、どんな方ですか？', 22, True, GOLD, PP_ALIGN.CENTER)
box(s, 3600000, 4300000, 4992000, 1650000, fill=WHITE)
text(s, 3850000, 4500000, 3200000, 400000, [('U-ライティング 上原 龍之介', 13, True, INK), ('サービスサイト', 10, False, SUB), ('https://www.u-promotion.com/', 11, True, GOLD)], spacing=1.5)
pic(s, f'{A}/qr.png', 7300000, 4450000, w=1150000)
notes(s, """
▼紹介を目指す場合：その場で名前を出してもらう
今日のお話で、思い浮かんだ方はいらっしゃいますか？
・出なければ「1週間後に一度ご連絡してもいいですか？」とフォローの約束を取る。

▼最後に必ず
逆に、私から○○さんにご紹介できるのはどんな方ですか？

▼打ち合わせ後（当日中）
・お礼の連絡と、決まった次の行動を文字で送る。
・相手から聞いた「紹介してほしい人」を控え、こちらから先に紹介できる人を探す。
""")

prs.save(OUT)
print('saved', OUT, 'slides:', len(prs.slides))
