import copy
import sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

SRC, OUT, QR = sys.argv[1], sys.argv[2], sys.argv[3]
prs = Presentation(SRC)
S = list(prs.slides)  # original order, 0-based


def set_tf(tf, paras):
    """Replace text keeping the first run's formatting of each paragraph."""
    ps = list(tf.paragraphs)
    while len(ps) < len(paras):
        new = copy.deepcopy(ps[-1]._p)
        ps[-1]._p.addnext(new)
        ps = list(tf.paragraphs)
    for p in ps[len(paras):]:
        p._p.getparent().remove(p._p)
    for p, text in zip(tf.paragraphs, paras):
        runs = p.runs
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        if runs:
            runs[0].text = text
        else:
            p.add_run().text = text


def sh(slide, idx):
    return list(slide.shapes)[idx]


def T(slide, idx, *paras):
    set_tf(sh(slide, idx).text_frame, list(paras))


def runs(slide, idx, texts):
    """Set texts run by run in the first paragraph (keeps mixed colors)."""
    rs = sh(slide, idx).text_frame.paragraphs[0].runs
    for r, t in zip(rs, texts):
        r.text = t
    for r in rs[len(texts):]:
        r._r.getparent().remove(r._r)


def delete(shape):
    shape._element.getparent().remove(shape._element)


def duplicate(slide):
    layout = slide.slide_layout
    new = prs.slides.add_slide(layout)
    for s in list(new.shapes):
        delete(s)
    for el in slide.shapes._spTree:
        if el.tag.endswith(('}sp', '}pic', '}grpSp', '}graphicFrame', '}cxnSp')):
            new.shapes._spTree.append(copy.deepcopy(el))
    rid_map = {}
    for rel in slide.part.rels.values():
        if 'notesSlide' in rel.reltype or 'slideLayout' in rel.reltype:
            continue
        rid_map[rel.rId] = new.part.relate_to(rel._target, rel.reltype)
    xml = new.shapes._spTree
    for el in xml.iter():
        for k, v in list(el.attrib.items()):
            if k.endswith('}embed') or k.endswith('}link') or k.endswith('}id'):
                if v in rid_map:
                    el.set(k, rid_map[v])
    return new


def label(slide, pic_idx, text):
    """Replace a gold '課題 0X' tag picture with an editable tag."""
    pic = sh(slide, pic_idx)
    l, t, w, h = pic.left, pic.top, pic.width, pic.height
    delete(pic)
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = GOLD
    box.line.fill.background()
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.name = 'Noto Sans JP'
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


GOLD = RGBColor(0xC6, 0xA1, 0x5B)


def tag(slide, pic_idx, text, width):
    """Replace the dark category tag picture with editable text."""
    pic = sh(slide, pic_idx)
    l, t, h = pic.left, pic.top, pic.height
    delete(pic)
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, Emu(width), h)
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0x2B, 0x2B, 0x2B)
    box.line.fill.background()
    tf = box.text_frame
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.name = 'Noto Sans JP'
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


# ---- 2: INDEX -> 本日の流れ
s = S[1]
T(s, 3, 'AGENDA')
sh(s, 3).width = Emu(4000000)
runs(s, 4, ['01', ' ', '○○様のお仕事について'])
runs(s, 6, ['02 ', '私の仕事と想い'])
runs(s, 8, ['03 ', 'お役に立てること・事例'])
runs(s, 10, ['04', ' ', '次の一歩'])
shapes = list(s.shapes)
for x in shapes[12:22] + shapes[26:31]:  # items 05-09 and their rules
    delete(x)

# ---- 3: プロフィール
s = S[2]
T(s, 13, 'ゲームアプリのメールサポート業務（2年）に従事')
T(s, 16, '建築積算の業務（3.5年）を経験')
T(s, 19, 'ライター歴6年以上。集客につながる文章・コンテンツ制作を中心に対応')
T(s, 9, '仕事の理念：マーケティング迷子を減らす')
T(s, 10, '本当に必要なものだけを提案し、誇大広告は絶対にしません。',
  '消費者とクライアントが笑顔でつながる機会を増やすことが私の信念です。')

# ---- 4: 課題3選 -> 理念の経緯（複製して紹介キーワード用にも使う）
keywords = duplicate(S[3])
s = S[3]
T(s, 6, 'WHY I DO THIS')
T(s, 7, '「マーケティング迷子」を減らしたい')
sh(s, 7).width = Emu(10668000)
T(s, 11, '高いお金をかけたのに成果が出ない')
T(s, 12, '数十万円かけてHP・LPを作った、講座やコンサルを受けた。それでも問い合わせが増えない。')
T(s, 14, '自力で頑張るが結果まで遠回りに')
T(s, 15, 'もう一度自分で集客に取り組むものの、手探りのため結果が出るまで時間がかかる。')
T(s, 17, '専門家に頼みたいが費用が重い')
T(s, 18, '早く結果を出すには専門家の力が必要。でも、また大きなお金がかかってしまう。')
T(s, 8, '広告は「投資」。だから、本当に必要なものだけを提案・制作します。')
for i, n in ((16, '声 03'), (13, '声 02'), (10, '声 01')):
    label(s, i, n)

s = keywords
T(s, 6, 'REFERRAL KEYWORDS')
T(s, 7, 'こんな声を聞いたら、ご紹介ください')
sh(s, 7).width = Emu(10668000)
T(s, 11, '「チラシやLPを作ったけど反応がない」')
T(s, 12, '作った後に成果が出ていない方。原因の診断から改善までお手伝いします。')
T(s, 14, '「この表現、薬機法的に大丈夫？」')
T(s, 15, '美容・健康分野で表現に悩む方。守りながら「売れる」表現を作ります。')
T(s, 17, '「集客まで手が回らない」')
T(s, 18, '本業が忙しい方、迷っている方。制作代行と集客顧問で支えます。')
T(s, 8, '合言葉は「マーケティング迷子」です。')
for i, n in ((16, 'KEYWORD 03'), (13, 'KEYWORD 02'), (10, 'KEYWORD 01')):
    label(s, i, n)

# ---- 5: サービス概要（紹介料 -> 価格）
s = S[4]
T(s, 15, '何から手をつけるべきか迷っている方向け')
T(s, 34, '料金：')
T(s, 36, '単発')
T(s, 37, '5,000円〜')
T(s, 33, '※30秒のスピーチ原稿の場合。LP制作は100,000円〜。')
T(s, 39, '料金：')
T(s, 40, '月額')
T(s, 41, '49,800円〜')
for i, left in ((37, 3250000), (41, 8850000)):
    x = sh(s, i)
    x.left, x.width = Emu(left), Emu(2650000)
    x.text_frame.paragraphs[0].runs[0].font.size = Pt(34)

    x.top = x.top + Emu(90000)
T(s, 38, '※月4時間まで。追加タスクは別途お見積り。')

# ---- 6: 選ばれる理由を3列に
s = S[5]
L, R, GAP = 1000125, 11182800, 300000
W = (R - L - 2 * GAP) // 3
heads, rules, bgs, bodies = [6, 8], [9, 10], [7, 11], [12, 13]
for col, (h, ru, bg, bd) in enumerate(zip(heads, rules, bgs, bodies)):
    for i in (h, ru, bg, bd):
        x = sh(s, i)
        x.left, x.width = L + col * (W + GAP), W
tree = s.shapes._spTree
new_els = [copy.deepcopy(sh(s, i)._element) for i in (8, 10, 11, 13)]
for el in new_els:
    tree.append(el)
n = len(list(s.shapes))
new = list(s.shapes)[n - 4:]
for x in new:
    x.left, x.width = L + 2 * (W + GAP), W
T(s, 6, '❶ 守りと攻めの薬機法対応')
T(s, 8, '❷ 集客の伴走支援')
T(s, n - 4, '❸ 必要なものだけを提案')
set_tf(new[3].text_frame, ['ご相談を受けても、本当に必要かを判断した上で受注します。不要なものは「作らなくていい」と正直にお伝えします。'])
for x in (sh(s, 12), sh(s, 13), new[3]):
    x.height = Emu(1300000)
for x in (sh(s, 7), sh(s, 11), new[2]):
    x.height = Emu(1750000)

# ---- お客様の声（選ばれる理由を複製）
voices = duplicate(s)
v = voices
T(v, 4, 'VOICE & WORKS')
T(v, 5, 'お客様の声・その他の実績')
T(v, 6, '提案力の高さ')
T(v, 8, 'ヒアリング力の高さ')
T(v, n - 4, '薬機法も対策可')
T(v, 12, '「こちらの要望も丁寧に聞いてくださり、的確な提案をいただけました。リーズナブルな価格にもかかわらず、しっかりとした内容でとても助かりました。」（健康装具の販売者 Sさん）')
T(v, 13, '「これまでの広告課題やお客様の反応を丁寧にヒアリングしてくださり、納得感のある広告に仕上げていただけました。」（脱毛・ホワイトニングサロン Yさん）')
T(v, n - 1, '「薬機法が関わる案件でも、安心して上原さんにすべてお任せできました。丁寧な対応と、高い完成度の原稿にとても満足しています。」（デザイン制作会社 Hさん）')
for i in (12, 13, n - 1):
    sh(v, i).height = Emu(1500000)
for i in (7, 11, n - 2):
    sh(v, i).height = Emu(1650000)
bar = v.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(1000125), Emu(5420000), Emu(10182675), Emu(850000))
bar.fill.solid()
bar.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
bar.line.color.rgb = GOLD
tf = bar.text_frame
tf.margin_left = Emu(250000)
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
for k, (txt, size, bold, color) in enumerate((
        ('その他の実績11件・過去の成果物はWebサイトで公開中', 14, True, RGBColor(0x33, 0x33, 0x33)),
        ('https://www.u-promotion.com/', 12, False, GOLD))):
    p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = txt
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = 'Noto Sans JP'
    r.font.color.rgb = color
v.shapes.add_picture(QR, Emu(10350000), Emu(5470000), Emu(750000), Emu(750000))

# ---- 7, 8, 9: 事例
T(S[6], 9, '1ヶ月で83.9万円', '目標比279%を達成')
T(S[7], 15, '3日で約50万円', '最終219万円を達成')
T(S[8], 13, '月平均50万円から', '最大150万円に改善')
tag(S[6], 6, 'クラファン｜ページ企画・ライティング担当', 3300000)
tag(S[7], 8, 'クラファン｜ページ企画・ライティング担当', 3300000)

# ---- 11: 料金一覧
s = S[10]
tbl = [x for x in s.shapes if x.has_table][0].table
set_tf(tbl.cell(4, 0).text_frame, ['集客顧問「ミチスジ」', '提案プラン'])
set_tf(tbl.cell(3, 0).text_frame, ['集客顧問「ミチスジ」', '相談プラン'])
set_tf(tbl.cell(3, 1).text_frame, ['「ミチスジ」の体験版。最大2時間のZoom相談で、', '現在の広告物を添削、対応方針を提案します。'])
tbl.cell(3, 2).text_frame.paragraphs[0].runs[-1].text = '14,800円〜'
T(s, 11, '「ミチスジ」顧問プラン契約期間中は、制作代行「カワリミ」の費用が 20% OFF になります。')

# ---- 12: カワリミ詳細
T(S[11], 16, '企画・構成から制作まで丸投げOK。 本業に集中できます。',
  '※制作に必要な情報・素材のご提供のみお願いします。')

# ---- 13: 制作メニュー
tbl = [x for x in S[12].shapes if x.has_table][0].table
c = tbl.cell(0, 2).text_frame.paragraphs[0].runs
c[-1].text = '）'

# ---- 14: ミチスジ詳細
T(S[13], 11, '月額 49,800円〜')
T(S[13], 12, '（月最大4時間・制作代行 20%OFF特典付き）')

# ---- 15: 体験プラン（次の一歩用に複製）
nxt = duplicate(S[14])
T(S[14], 11, '単発 14,800円〜')

s = nxt
T(s, 5, 'NEXT STEP')
T(s, 6, '次の一歩')
T(s, 8, '本日のお話をもとに、', '必要なものだけをご提案します。')
T(s, 9, 'ご状況に合わせて3つの進め方からお選びいただけます。必要のない契約はおすすめしません。')
T(s, 10, '進め方')
T(s, 11, 'ご提案・お見積りは無料です')
T(s, 21, '① ご提案・お見積り（無料）')
T(s, 15, '今日伺った課題に対し、制作・顧問の提案と見積りをお送りします。')
T(s, 22, '② 相談プラン（14,800円〜）')
T(s, 16, 'まだ迷っている方は、最大2時間のZoom相談で現状を一緒に整理します。')
T(s, 23, '③ お知り合いのご紹介')
T(s, 17, '集客や広告でお困りの方がいれば、ぜひおつなぎください。')
for i in (21, 22, 23):
    sh(s, i).width = Emu(4000000)

# ---- 16: 提案プラン
T(S[15], 11, '単発 49,800円〜')
T(S[15], 9, 'ヒアリングと現状データをもとに分析し、必要な施策をまとめた「提案書」を納品します。',
  '※分析範囲により変動します。')

# ---- 18: 紹介パートナー制度
for para in sh(S[17], 17).text_frame.paragraphs:
    for r in para.runs:
        r.text = r.text.replace('20万円', '10万円')

# ---- 並べ替え
order = [0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 'voices', 10, 'next', 'kw', 16, 17, 11, 12, 13, 14, 15]
ids = list(prs.slides._sldIdLst)
by_part = {prs.slides.get(int(i.get('id'))).part: i for i in ids} if False else None
lst = prs.slides._sldIdLst
els = list(lst)
orig_els, nxt_el, kw_el, v_el = els[:18], None, None, None
for el in els[18:]:
    slide = prs.slides.get(int(el.get('id')))
    if slide is nxt:
        nxt_el = el
    elif slide is keywords:
        kw_el = el
    elif slide is voices:
        v_el = el
for el in els:
    lst.remove(el)
for o in order:
    lst.append({'next': nxt_el, 'kw': kw_el, 'voices': v_el}.get(o, None) if isinstance(o, str) else orig_els[o])

prs.save(OUT)
print('saved', OUT)
